import re
import io
import logging
from typing import List, Dict, Any, Tuple
from app.core.config import settings

logger = logging.getLogger("mailscope.extractor")

# Robust email pattern
EMAIL_REGEX = re.compile(r'[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}')

class DocumentExtractorService:
    def extract_text_from_pdf(self, file_bytes: bytes) -> str:
        """Extract text from a PDF document using pypdf."""
        try:
            import pypdf
            reader = pypdf.PdfReader(io.BytesIO(file_bytes))
            text_parts = []
            for page in reader.pages:
                txt = page.extract_text()
                if txt:
                    text_parts.append(txt)
            return "\n".join(text_parts)
        except Exception as e:
            logger.error(f"Error parsing PDF: {e}")
            raise ValueError(f"Failed to extract text from PDF document: {str(e)}")

    def extract_text_from_docx(self, file_bytes: bytes) -> str:
        """Extract text from a Word DOCX document using python-docx."""
        try:
            import docx
            doc = docx.Document(io.BytesIO(file_bytes))
            text_parts = []
            for para in doc.paragraphs:
                if para.text:
                    text_parts.append(para.text)
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        if cell.text:
                            text_parts.append(cell.text)
            return "\n".join(text_parts)
        except Exception as e:
            logger.error(f"Error parsing DOCX: {e}")
            raise ValueError(f"Failed to extract text from DOCX document: {str(e)}")

    def extract_text_from_doc(self, file_bytes: bytes) -> str:
        """Fallback text extraction for legacy binary .doc format."""
        try:
            # Basic string decode fallback for binary .doc files
            raw_text = file_bytes.decode("ascii", errors="ignore")
            matches = EMAIL_REGEX.findall(raw_text)
            if matches:
                return " ".join(matches)
            raise ValueError("Legacy binary .doc file format could not be decoded cleanly. Please convert file to .docx or .pdf.")
        except Exception as e:
            raise ValueError("Legacy binary .doc files are not natively supported. Please convert your file to .docx or .pdf.")

    def clean_email(self, match: str) -> str:
        """Clean surrounding punctuation and normalize email address."""
        cleaned = match.strip(".,;:()[]\"'<>?!")
        return cleaned.strip()

    def process_document(self, filename: str, file_bytes: bytes) -> Dict[str, Any]:
        """Process document file and extract email addresses with deduplication stats."""
        ext = filename.lower().split(".")[-1]
        
        if ext == "pdf":
            raw_text = self.extract_text_from_pdf(file_bytes)
        elif ext == "docx":
            raw_text = self.extract_text_from_docx(file_bytes)
        elif ext == "doc":
            raw_text = self.extract_text_from_doc(file_bytes)
        else:
            raise ValueError(f"Unsupported file format '.{ext}'. Supported formats: .pdf, .docx, .doc")

        if not raw_text or not raw_text.strip():
            raise ValueError("No text content could be extracted from the uploaded document.")

        raw_matches = EMAIL_REGEX.findall(raw_text)
        cleaned_matches = []
        for m in raw_matches:
            c = self.clean_email(m)
            if c and "@" in c and "." in c.split("@")[-1]:
                cleaned_matches.append(c)

        total_found = len(cleaned_matches)
        if total_found == 0:
            raise ValueError("No email addresses were found in the uploaded document.")

        # Case-insensitive deduplication while retaining original casing preference
        unique_map = {}
        for email in cleaned_matches:
            key = email.lower()
            if key not in unique_map:
                unique_map[key] = email

        unique_emails = list(unique_map.values())
        unique_count = len(unique_emails)
        duplicates_count = total_found - unique_count

        max_allowed = settings.MAX_BULK_EMAILS
        if unique_count > max_allowed:
            raise ValueError(f"Too many email addresses found in document ({unique_count}). Maximum allowed: {max_allowed}.")

        return {
            "filename": filename,
            "total_found": total_found,
            "unique_count": unique_count,
            "duplicates_count": duplicates_count,
            "emails": unique_emails
        }

document_extractor = DocumentExtractorService()
